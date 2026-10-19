from django.http import HttpResponseRedirect, HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.contrib.auth import authenticate, login, get_user_model
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth.forms import UserCreationForm
from django.contrib.auth.views import PasswordResetView
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from datetime import timedelta
from django.conf import settings
import openpyxl
import csv
from .models import DonationRequest
from openpyxl import load_workbook
import os

from .models import CustomUser, Delivery, DonationRequest, Item, Message
from .forms import CustomUserCreationForm, QueryForm
from .utils import send_sms

User = get_user_model()


# ✅ KEEP THIS ONE — only version of delete_item
@login_required
def delete_item(request, item_id):
    item = get_object_or_404(Item, id=item_id)
    if request.method == 'POST':
        item_name = item.name
        item.delete()
        messages.success(request, f'Item "{item_name}" has been deleted successfully!')
        return redirect('my_items')
    else:
        return render(request, 'delete_item.html', {'item': item})


def register(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')
        password = request.POST.get('password')
        confirm_password = request.POST.get('confirm_password')
        full_name = request.POST.get('full_name')
        phone_number = request.POST.get('phone_number')
        alt_phone_number = request.POST.get('alt_phone_number')
        address = request.POST.get('address')
        gender = request.POST.get('gender')

        if password != confirm_password:
            messages.error(request, "Passwords do not match.")
            return redirect('register')

        if CustomUser.objects.filter(username=username).exists():
            messages.error(request, "Username already taken.")
            return redirect('register')

        if CustomUser.objects.filter(email=email).exists():
            messages.error(request, "Email already registered.")
            return redirect('register')

        CustomUser.objects.create_user(
            username=username,
            email=email,
            password=password,
            full_name=full_name,
            phone_number=phone_number,
            alternative_phone_number=alt_phone_number,
            address=address,
            gender=gender,
            theme='light',
            nickname=''
        )

        messages.success(request, "Registration successful. You can now log in.")
        return redirect('login')

    return render(request, 'registration/register.html')


def item_detail(request, pk):
    item = get_object_or_404(Item, pk=pk)
    return render(request, 'item_detail.html', {'item': item})


@login_required
def login_redirect(request):
    if request.user.is_superuser:
        return redirect('admin_dashboard')
    else:
        return redirect('user_dashboard')


@login_required
def user_dashboard(request):
    items = Item.objects.filter(approved=True).order_by('-created_at')
    total_items = Item.objects.filter(approved=True).count()
    trade_items = Item.objects.filter(approved=True, category='trade').count()
    donation_items = Item.objects.filter(approved=True, category='donate').count()
    my_items = Item.objects.filter(owner=request.user).count()

    trade_percentage = (trade_items / total_items * 100) if total_items > 0 else 0
    donation_percentage = (donation_items / total_items * 100) if total_items > 0 else 0
    my_percentage = (my_items / total_items * 100) if total_items > 0 else 0

    return render(request, "dashboard/user_dashboard.html", {
        "items": items,
        "total_items": total_items,
        "trade_items": trade_items,
        "donation_items": donation_items,
        "my_items": my_items,
        "trade_percentage": trade_percentage,
        "donation_percentage": donation_percentage,
        "my_percentage": my_percentage,
    })


@login_required
@user_passes_test(lambda u: u.is_superuser)
def admin_dashboard(request):
    one_week_ago = timezone.now() - timedelta(days=7)
    new_users = User.objects.filter(date_joined__gte=one_week_ago).count()
    total_users = User.objects.count()
    total_items = Item.objects.count()
    pending_items = Item.objects.filter(approved=False).count()
    all_users = User.objects.all()
    all_items = Item.objects.all()
   
    # Fetch all donation requests for the admin report
    donation_requests = DonationRequest.objects.all().order_by('-created_at')

    context = {
        'total_users': total_users,
        'new_users': new_users,
        'total_items': total_items,
        'pending_items': pending_items,
        'all_users': all_users,
        'all_items': all_items,
        'donation_requests': donation_requests,
    }

    return render(request, 'dashboard/admin_dashboard.html', context)


def item_list(request):
    items = Item.objects.all()
    return render(request, 'item_list.html', {'items': items})


@staff_member_required
def manage_items(request):
    filter_type = request.GET.get('filter', None)
    if filter_type == 'pending':
        items = Item.objects.filter(approved=False).order_by('-created_at')
    else:
        items = Item.objects.all().order_by('-created_at')
    return render(request, "admin/manage_items.html", {"items": items, "filter_type": filter_type})


@staff_member_required
def manage_user(request):
    filter_type = request.GET.get('filter', None)
    if filter_type == 'new':
        one_week_ago = timezone.now() - timedelta(days=7)
        users = User.objects.filter(date_joined__gte=one_week_ago)
    else:
        users = User.objects.all()

    context = {
        'users': users,
        'filter_type': filter_type,
        'user': request.user
    }
    return render(request, 'admin/manage_user.html', context)


@staff_member_required
def approve_item(request, item_id):
    item = get_object_or_404(Item, id=item_id)
    item.approved = True
    item.save()
    messages.success(request, f"Item '{item.name}' has been approved ✅")
    return redirect("manage_items")


@user_passes_test(lambda u: u.is_superuser)
def schedule_delivery(request, item_id):
    item = get_object_or_404(Item, id=item_id)
    if request.method == 'POST':
        delivery_address = request.POST['delivery_address']
        delivery_date = request.POST['delivery_date']
        Delivery.objects.create(
            item=item,
            delivery_address=delivery_address,
            delivery_date=delivery_date,
            delivered_by=request.user
        )
        return redirect('admin_dashboard')
    return render(request, 'schedule_delivery.html', {'item': item})


def login_view(request):
    if request.method == "POST":
        username = request.POST.get("username")
        password = request.POST.get("password")
        role = request.POST.get("role")
        user = authenticate(request, username=username, password=password)

        if user is not None:
            if role == "admin" and user.is_superuser:
                login(request, user)
                return redirect('admin_dashboard')
            elif role == "user" and not user.is_superuser:
                login(request, user)
                return redirect('user_dashboard')
            else:
                messages.error(request, "Invalid role selection for this account.")
        else:
            messages.error(request, "Invalid username or password.")

    return render(request, "login.html")


@login_required
def settings_view(request):
    if request.method == 'POST':
        user = request.user
        user.address = request.POST.get('address')
        user.nickname = request.POST.get('nickname')
        new_password = request.POST.get('password')
        if new_password and new_password != "•••••••••••••••":
            user.set_password(new_password)
        theme = request.POST.get('theme')
        if theme:
            user.theme = theme
        user.save()
        messages.success(request, "Settings updated successfully!")
        return redirect('settings')

    return render(request, 'admin/settings.html', {'user': request.user})


class CustomPasswordResetView(PasswordResetView):
    template_name = 'forgot_pass.html'
    email_template_name = 'registration/password_reset_email.html'
    success_url = reverse_lazy('password_reset_done')


def forgot_pass(request):
    return CustomPasswordResetView.as_view()(request)


def user_details(request, id):
    user = get_object_or_404(User, id=id)
    return render(request, 'admin/user_details.html', {'user': user})


def add_user(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')
        password = request.POST.get('password')
        User.objects.create_user(username=username, email=email, password=password)
        messages.success(request, "User added successfully.")
        return redirect('manage_user')
    return render(request, 'admin/add_user.html')


def edit_user(request, user_id):
    user = get_object_or_404(User, id=user_id)
    if request.method == "POST":
        user.username = request.POST.get('username')
        user.email = request.POST.get('email')
        user.save()
        messages.success(request, "User updated successfully.")
        return redirect('manage_user')
    return render(request, 'admin/edit_user.html', {'user': user})


def delete_user(request, user_id):
    user = get_object_or_404(User, id=user_id)
    if request.method == "POST":
        user.delete()
        messages.success(request, "User deleted successfully.")
        return redirect('manage_user')
    return render(request, 'admin/delete_user.html', {'user': user})


def add_item(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        category = request.POST.get('category')
        description = request.POST.get('description')
        image = request.FILES.get('image')
        Item.objects.create(
            name=name,
            category=category,
            description=description,
            image=image,
            owner=request.user,
            approved=False
        )
        messages.success(request, "Item submitted successfully. Awaiting admin approval.")
        return redirect('user_dashboard')

    return render(request, 'add_item.html')


def user_add_item(request):
    if request.method == 'POST':
        name = request.POST.get('name')
        category = request.POST.get('category')
        description = request.POST.get('description')
        image = request.FILES.get('image')
        Item.objects.create(
            name=name,
            category=category,
            description=description,
            image=image,
            owner=request.user,
            approved=False
        )
        messages.success(request, "Item submitted successfully. Awaiting admin approval.")
        return redirect('user_add_item')

    return render(request, 'user/user_add_item.html')


@login_required
def message_view(request, user_id=None):
    chats = []
    messages_list = []
    active_chat_user = None

    if user_id:
        active_chat_user = User.objects.get(id=user_id)
        messages_list = Message.objects.filter(
            sender__in=[request.user, active_chat_user],
            receiver__in=[request.user, active_chat_user]
        ).order_by('timestamp')
        Message.objects.filter(receiver=request.user, sender=active_chat_user).update(is_read=True)

    users = User.objects.exclude(id=request.user.id)
    for user in users:
        last_msg = Message.objects.filter(
            sender__in=[request.user, user],
            receiver__in=[request.user, user]
        ).order_by('-timestamp').first()
        unread_count = Message.objects.filter(receiver=request.user, sender=user, is_read=False).count()
        chats.append({
            'user': user,
            'last_message': last_msg.content if last_msg else '',
            'last_message_time': last_msg.timestamp if last_msg else '',
            'unread_count': unread_count
        })

    if request.method == "POST":
        content = request.POST.get("content")
        Message.objects.create(sender=request.user, receiver=active_chat_user, content=content)
        send_sms(active_chat_user.phone_number, content)
        return redirect('chat_user', user_id=active_chat_user.id)

    return render(request, 'chat_user.html', {
        'chats': chats,
        'messages': messages_list,
        'active_chat_user': active_chat_user
    })


def search_users(request):
    query = request.GET.get('q', '')
    users = User.objects.filter(username__icontains=query).exclude(id=request.user.id)[:5]
    data = [{'username': u.username, 'phone_number': getattr(u, 'phone_number', ''), 'id': u.id} for u in users]
    return JsonResponse(data, safe=False)


def donate_item(request, item_id):
    item = get_object_or_404(Item, id=item_id)
    if request.method == "POST":
        return HttpResponseRedirect(reverse('item_list'))
    return render(request, 'donation.html', {'item': item})


def trade_item(request, item_id):
    item = get_object_or_404(Item, id=item_id)
    if request.method == "POST":
        return HttpResponseRedirect(reverse('item_list'))
    return render(request, 'trading.html', {'item': item})


@login_required
def profile_view(request):
    return render(request, 'admin/profile.html')


@login_required
def user_profile(request):
    user = request.user
    if request.method == 'POST':
        email = request.POST.get('email')
        if not email:
            print("Email field is missing or empty")
            return redirect('user_profile')
        user.full_name = request.POST.get('full_name')
        user.nickname = request.POST.get('nickname')
        user.email = email
        user.phone = request.POST.get('phone')
        user.bio = request.POST.get('bio')
        user.save()
    return render(request, 'user/user_profile.html', {'user': user})


def welcome(request):
    return render(request, 'welcome.html')


def my_items(request):
    items = Item.objects.filter(owner=request.user).order_by('-created_at')
    return render(request, 'user/my_items.html', {'items': items})


@login_required
def edit_item(request, item_id):
    item = get_object_or_404(Item, id=item_id, owner=request.user)
    if request.method == 'POST':
        form = ItemForm(request.POST, request.FILES, instance=item)
        if form.is_valid():
            form.save()
            messages.success(request, 'Item updated successfully!')
            return redirect('user/my_items.html')
    else:
        form = ItemForm(instance=item)
    return render(request, 'edit_item.html', {'form': form, 'item': item})


def edit_profile(request):
    if request.method == 'POST':
        user = request.user
        user.username = request.POST.get('username')
        user.email = request.POST.get('email')
        user.save()
        messages.success(request, "Profile updated successfully!")
        return redirect('user_profile')
    return render(request, 'user/edit_profile.html')


def submit_query(request):
    if request.method == 'POST':
        form = QueryForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Your query has been submitted successfully!')
            return redirect('home')
    else:
        form = QueryForm()
    return render(request, 'home.html', {'form': form})


def trade_item_view(request):
    return render(request, 'trading.html')


def donation_view(request, item_id):
    item = get_object_or_404(Item, id=item_id)

    if request.method == 'POST':
        donation_request = DonationRequest.objects.create(
            item=item,
            requester=request.user,
            reason=request.POST.get('reason'),
            organization=request.POST.get('organization'),
            full_name=request.POST.get('full_name'),
            phone=request.POST.get('phone'),
            email=request.POST.get('email'),
            address=request.POST.get('address'),
            city=request.POST.get('city'),
            postal_code=request.POST.get('postal_code'),
        )

        messages.success(request, "Donation request submitted successfully! Await admin approval.")
        return render(request, 'donation.html', {'item': item})

    return render(request, 'donation.html', {'item': item})


def logistic_view(request):
    if request.method == 'POST':
        driver = request.POST.get('driver')
        customer_name = request.POST.get('customer_name')
        address = request.POST.get('address')
        contact_number = request.POST.get('contact_number')
        delivery_items = request.POST.get('delivery_items')
        pickup_items = request.POST.get('pickup_items')
        delivery_date = request.POST.get('delivery_date')
        delivery_time = request.POST.get('delivery_time')
        location = request.POST.get('location')

        Delivery.objects.create(
            item=None,
            driver=driver,
            customer_name=customer_name,
            address=address,
            contact_number=contact_number,
            delivery_items=delivery_items,
            pickup_items=pickup_items,
            delivery_date=delivery_date,
            delivery_time=delivery_time,
            location=location,
            delivery_address=address,
            delivered_by=request.user,
        )

        file_path = os.path.join(settings.BASE_DIR, 'delivery_schedule.xlsx')
        if not os.path.exists(file_path):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Daily Delivery Schedule"
            ws.append([
                "Driver", "Customer Name", "Address", "Contact Number",
                "Delivery Item(s)", "Item(s) to Pick Up", "Location",
                "Date", "Time", "Status"
            ])
        else:
            wb = load_workbook(file_path)
            ws = wb.active

        ws.append([
            driver, customer_name, address, contact_number,
            delivery_items, pickup_items, location,
            delivery_date, delivery_time, "Pending"
        ])
        wb.save(file_path)

        messages.success(request, f"Delivery scheduled for {customer_name} on {delivery_date} at {delivery_time}")
        return redirect('admin_logistic')

    deliveries = Delivery.objects.all().order_by('-delivery_date', '-delivery_time')
    return render(request, 'admin/logistic.html', {'deliveries': deliveries})


def search_items(request):
    query = request.GET.get('q')
    items = Item.objects.filter(name__icontains=query) if query else []
    return render(request, 'search_results.html', {'items': items, 'query': query})


def filter_items(request):
    category = request.GET.get('category')
    items = Item.objects.filter(category=category) if category else []
    return render(request, 'filter_results.html', {'items': items, 'category': category})


def export_users_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Users Report"
    headers = ["Full Name", "Username", "Email", "Phone", "Address", "Joined", "Last Login", "Gender"]
    ws.append(headers)

    for user in CustomUser.objects.all():
        ws.append([
            user.full_name,
            user.username,
            user.email,
            user.phone_number,
            user.address,
            user.date_joined.strftime("%Y-%m-%d"),
            user.last_login.strftime("%Y-%m-%d %H:%M") if user.last_login else "",
            user.get_gender_display()
        ])

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = "attachment; filename=users_report.xlsx"
    wb.save(response)
    return response


def export_items_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Items Report"
    headers = ["Name", "Category", "Created At", "Owner"]
    ws.append(headers)

    for item in Item.objects.all():
        ws.append([
            item.name,
            item.category,
            item.created_at.strftime("%Y-%m-%d %H:%M"),
            item.owner.username if item.owner else ""
        ])

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = "attachment; filename=items_report.xlsx"
    wb.save(response)
    return response


@staff_member_required
def donation_requests(request):
   
      donation_requests = DonationRequest.objects.all().order_by('-created_at')
      return render(request, 'admin/donation_requests.html', {'donation_requests': donation_requests})


@login_required
def my_donations(request):
    user_requests = DonationRequest.objects.filter(requester=request.user).order_by('-created_at')
    return render(request, 'user/my_donations.html', {'donation_requests': user_requests})


def approve_donation(request, request_id):
    donation_request = get_object_or_404(DonationRequest, id=request_id)
    donation_request.status = 'approved'
    donation_request.save()
    messages.success(request, f"Donation request for {donation_request.item.name} approved!")
    
    # Optionally: Notify the requester and item owner here
    # send_notification(donation_request)
    
    return redirect('admin/donation_requests')


def decline_donation(request, request_id):
    donation_request = get_object_or_404(DonationRequest, id=request_id)
    donation_request.status = 'rejected'
    donation_request.save()
    messages.success(request, f"Donation request for {donation_request.item.name} declined!")
    
    # Optionally: Notify the requester here
    # send_notification(donation_request)
    
    return redirect('admin/donation_requests')


def export_donations(request):
    # Create the HttpResponse object with CSV header
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="donation_requests.csv"'

    writer = csv.writer(response)

    # Write CSV header
    writer.writerow([
        'Requester', 'Item', 'Item Owner', 'Organization', 'Reason',
        'Full Name', 'Phone', 'Email', 'Address', 'City', 'Postal Code',
        'Status', 'Scheduled Date', 'Scheduled Time', 'Created At'
    ])

    # Query all donation requests
    requests = DonationRequest.objects.all()

    for r in requests:
        writer.writerow([
            r.requester.username,
            r.item.name,
            r.item.owner.username,
            r.organization,
            r.reason,
            r.full_name,
            r.phone,
            r.email,
            r.address,
            r.city,
            r.postal_code,
            r.status,
            r.scheduled_date if r.scheduled_date else '',
            r.scheduled_time if r.scheduled_time else '',
            r.created_at.strftime("%Y-%m-%d %H:%M"),
        ])

    return response

